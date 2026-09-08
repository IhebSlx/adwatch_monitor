"""Die Übergabe an Daniel: Architekturbüros mit Tätigkeit in einem Zielland.

WAS DRIN STEHT UND WAS BEWUSST NICHT.
Jede Spalte ist entweder ein FAKT aus dem CRM oder ein am Website-Text
BELEGTER Fund. Kein Feld ist ein Modellurteil.

`solarlux_relevance` fehlt deshalb absichtlich. Die Note gäbe es zwar (der
Architekten-Prompt kann sie), aber sie wurde nie gegen ein Ergebnis geprüft:
gemessen 2026-09-08 war die Überschneidung zwischen bewerteten Büros und Büros
mit gewonnenem Objekt exakt null. Eine ungeprüfte Note in einer Übergabe sieht
aus wie eine Zahl und wird sortiert wie eine — das wäre schlechter als sie
wegzulassen. (Der Prüflauf dafür scheiterte am leeren Anthropic-Guthaben, nicht
am Verfahren.)

DIE SORTIERUNG IST DIE EIGENTLICHE AUSSAGE.
Zuerst die Beziehungsstufe, dann der gewonnene Wert. Ganz oben stehen also
Büros, mit denen Solarlux schon gebaut hat UND die im Zielland tätig sind —
und genau die stehen fast nie im Zielland selbst.
"""
import io
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from sqlalchemy import text                       # noqa: E402

from adwatch import config                        # noqa: E402
from adwatch.db import SessionLocal, init_db      # noqa: E402
from adwatch.insights.beziehung import STUFEN     # noqa: E402

SPALTEN = [
    ("stufe", "Beziehung (0-5)"), ("warum", "Beziehung — was genau"),
    ("name", "Büro"), ("stadt", "Sitz — Ort"), ("land", "Sitz — Land"),
    ("website", "Website"),
    ("zielstaedte", "Orte im Zielland"),
    ("taetig_in", "Tätig in (sicher)"),
    ("moeglich", "Tätig in (möglich)"),
    ("rolle", "Rolle"), ("rolle_beleg", "Rolle — Beleg"),
    ("objekte", "Objekte im CRM"), ("gewonnen", "davon gewonnen"),
    ("gewonnener_wert", "Gewonnener Wert (EUR)"),
    ("beleg", "Beleg für das Zielland"),
]


def zeilen(land: str):
    with SessionLocal() as s:
        roh = s.execute(text("""
            SELECT name, city, country, website_domain,
                   active_countries, active_countries_all, active_countries_evidence,
                   active_cities, decision_role, decision_role_evidence,
                   COALESCE(relation_level,0), relation_why,
                   COALESCE(arch_projects,0), COALESCE(arch_won,0),
                   COALESCE(arch_won_value,0)
            FROM companies
            WHERE segment = 'Architekten' AND sub_segment = 'Architekturbüro'
              AND duplicate_of IS NULL AND active_countries IS NOT NULL
        """)).all()

    aus = []
    for (name, stadt, sitz, web, ac, alle, ev, cities, rolle, rbeleg,
         stufe, warum, objekte, gewonnen, wert) in roh:
        laender = json.loads(ac or "[]")
        if land.upper() not in [x.upper() for x in laender]:
            continue
        alle_d = json.loads(alle or "{}")
        aus.append({
            "stufe": stufe,
            "warum": warum or STUFEN.get(stufe, ""),
            "name": name, "stadt": stadt or "", "land": sitz or "",
            "website": web or "",
            "zielstaedte": ", ".join((json.loads(cities or "{}")).get(land.upper(), [])),
            "taetig_in": ", ".join(laender),
            "moeglich": ", ".join(k for k, v in alle_d.items() if v == "moeglich"),
            "rolle": rolle or "", "rolle_beleg": rbeleg or "",
            "objekte": objekte, "gewonnen": gewonnen,
            "gewonnener_wert": round(wert, 2),
            "beleg": " · ".join((json.loads(ev or "{}")).get(land.upper(), [])),
        })
    aus.sort(key=lambda r: (-r["stufe"], -r["gewonnener_wert"], r["name"]))
    return aus


def schreiben(land: str, pfad: str) -> dict:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    daten = zeilen(land)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Architekturbüros {land.upper()}"

    kopf = Font(bold=True, color="FFFFFF")
    fuell = PatternFill("solid", fgColor="1F3864")
    for i, (_, label) in enumerate(SPALTEN, start=1):
        z = ws.cell(row=1, column=i, value=label)
        z.font, z.fill = kopf, fuell
        z.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    # Warme Zeilen hervorheben — sie sind der Grund für die Datei.
    warm = PatternFill("solid", fgColor="E2EFDA")
    for r, zeile in enumerate(daten, start=2):
        for c, (key, _) in enumerate(SPALTEN, start=1):
            zelle = ws.cell(row=r, column=c, value=zeile[key])
            if zeile["stufe"] >= 3:
                zelle.fill = warm
    breiten = {"name": 38, "warum": 28, "zielstaedte": 40, "beleg": 46,
               "website": 28, "rolle_beleg": 30, "stadt": 16}
    for i, (key, _) in enumerate(SPALTEN, start=1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = breiten.get(key, 15)

    wb.save(pfad)
    return {"zeilen": len(daten),
            "warm": sum(1 for z in daten if z["stufe"] >= 3),
            "fremd": sum(1 for z in daten if (z["land"] or "").upper() != land.upper()),
            "vergibt": sum(1 for z in daten if z["rolle"] == "vergibt Aufträge"),
            "pfad": pfad}


if __name__ == "__main__":
    init_db()
    land = (sys.argv[1] if len(sys.argv) > 1 else "ES").upper()
    ziel = os.path.join(str(config.ROOT), "output",
                        f"Architekturbueros_{land}.xlsx")
    os.makedirs(os.path.dirname(ziel), exist_ok=True)
    print(schreiben(land, ziel))
